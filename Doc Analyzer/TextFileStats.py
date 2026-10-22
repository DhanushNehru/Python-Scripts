from docx import Document

def analyze_file(filename):
    try:
        doc = Document(filename)
        text = "\n".join(
            paragraph.text for paragraph in doc.paragraphs
        )

        print(f"Lines (paragraphs): {len(doc.paragraphs)}")
        print(f"Words: {len(text.split())}")
        print(f"Characters: {len(text)}")

    except FileNotFoundError:
        print("File not found.")
    except Exception as error:
        print(f"Could not read Word file: {error}")


if __name__ == "__main__":
    filename = input("Enter Word file path (.docx): ").strip()
    analyze_file(filename)
